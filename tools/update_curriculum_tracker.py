#!/usr/bin/env python3
"""Synchronize the DOCX tracker with the approved fundamentals-first strategy."""

from __future__ import annotations

from pathlib import Path

import docx


ROOT = Path(__file__).resolve().parent.parent
TRACKER = ROOT / "metadata" / "curriculum-tracker.docx"
START = "CURRENT STRATEGY — FUNDAMENTALS FIRST"
END = "END OF CURRENT STRATEGY ADDENDUM"
ANCHOR = "How to Use This Tracker"


def remove_previous_addendum(document: docx.Document) -> None:
    removing = False
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text == START:
            removing = True
        if removing:
            paragraph._element.getparent().remove(paragraph._element)
        if text == END:
            removing = False


def add_before(anchor, text: str, style: str | None = None):
    paragraph = anchor.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    return paragraph


def main() -> None:
    document = docx.Document(TRACKER)
    remove_previous_addendum(document)

    purpose_table = document.tables[0]
    purpose_cell = purpose_table.cell(0, 0)
    purpose_cell.paragraphs[1].text = (
        "Build a connected mental model from physical computation through "
        "CPU architecture, C, operating systems, Linux kernel internals, "
        "drivers, debugging, and virtualization. The primary career outcome "
        "is Linux kernel/platform development; distributed, cloud, and AI "
        "infrastructure remain later context where they depend on kernel "
        "mechanisms."
    )
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("From transistor to cloud"):
            paragraph.text = (
                "From transistor to Linux kernel—without treating any layer "
                "as magic."
            )
            break

    anchor = next(
        (p for p in document.paragraphs if p.text.strip() == ANCHOR),
        None,
    )
    if anchor is None:
        raise SystemExit(f"could not find tracker anchor: {ANCHOR!r}")

    add_before(anchor, START, "Heading 1")
    add_before(
        anchor,
        "Updated 30 September 2026. The detailed 66-week curriculum remains "
        "below, but the active learning sequence is controlled by confidence "
        "gates rather than elapsed calendar time.",
    )
    add_before(
        anchor,
        "Primary outcome: become a Linux kernel/platform developer with deep "
        "C, systems-programming, debugging, concurrency, memory-management, "
        "driver, virtualization, and upstream-development skill.",
    )

    add_before(anchor, "Active focus", "Heading 2")
    add_before(
        anchor,
        "Work through the prepared foundations in order. Keep one active "
        "module at a time. Prepared future PDFs do not imply progress or "
        "mastery.",
        "List Bullet",
    )
    add_before(
        anchor,
        "Specialization, market-validation sprints, active upstream issue "
        "selection, and ToyOS expansion remain deferred until their foundation "
        "gates are met.",
        "List Bullet",
    )

    add_before(anchor, "Foundation Gate A — CPU, C, assembly, toolchain", "Heading 2")
    for item in (
        "Trace a simple C expression through compiler output, ISA-visible "
        "state, and CPU execution.",
        "Read basic x86-64 assembly, ABI/stack frames, object files, linking, "
        "ELF loading, cache/TLB/NUMA, and C object/lifetime behavior.",
        "Write and debug warning-clean multi-file C with pointers, aggregates, "
        "allocation, cleanup, sanitizers, GDB, and edge-case tests.",
    ):
        add_before(anchor, item, "List Bullet")

    add_before(anchor, "Foundation Gate B — hardware/software boundary", "Heading 2")
    for item in (
        "Explain transistor/gate/timing behavior using bounded simulation "
        "evidence without turning the course into an HDL/ASIC specialty.",
        "Trace a teaching RISC-V core cycle by cycle and distinguish ISA, "
        "microarchitecture, OS policy, and physical behavior.",
        "Connect a C task to instructions, datapath/control, pipeline hazards, "
        "cache/TLB/MMU activity, gates, and captured state.",
    ):
        add_before(anchor, item, "List Bullet")

    add_before(anchor, "Foundation Gate C — operating-system mental model", "Heading 2")
    for item in (
        "Trace syscall, exception, interrupt, scheduling, page-fault, and "
        "return paths with context, state, locking, and lifetime annotations.",
        "Implement and debug simplified allocation, paging, context switching, "
        "syscall, descriptor, and IPC mechanisms in ToyOS or bounded models.",
        "Compare mechanisms across ToyOS, xv6, and Linux and diagnose at least "
        "one injected failure from symptom to regression test.",
    ):
        add_before(anchor, item, "List Bullet")

    add_before(anchor, "Required continuous lanes", "Heading 2")
    for item in (
        "C practice: strict warnings, tests, GDB, sanitizers, Valgrind, "
        "toolchain/assembly inspection, and systems-oriented projects.",
        "Debugging: user-space and kernel diagnosis using evidence, including "
        "perf/PMUs, ftrace/trace-cmd, eBPF, sanitizers, lockdep, kdump/crash, "
        "and QEMU/KGDB where appropriate.",
        "Core-kernel atlas: syscall, scheduling, page-fault, interrupt, I/O, "
        "network, driver, and KVM execution paths.",
        "Interview readiness: concise explanations, whiteboard traces, C "
        "coding, debugging scenarios, trade-offs, and project deep dives.",
        "Cross-layer systems depth: memory ordering/coherence, NUMA, PCIe, "
        "DMA/IOMMU/RDMA, scheduler/cgroups/namespaces, storage, networking, "
        "virtualization/confidential computing, and observability.",
    ):
        add_before(anchor, item, "List Bullet")

    add_before(anchor, "Later pivot and contribution gate", "Heading 2")
    add_before(
        anchor,
        "Default differentiator is KVM/x86/AMD virtualization and SEV-SNP "
        "inside a broader kernel/platform profile. Networking/eBPF, "
        "embedded/BSP, PCIe/accelerator, and storage paths are validated with "
        "bounded market and proof-of-fit sprints only after foundations.",
        "List Bullet",
    )
    add_before(
        anchor,
        "Required career evidence includes public upstream review and at least "
        "one meaningful, reproduced, tested kernel/KVM contribution; merge is "
        "preferred but not fully controllable.",
        "List Bullet",
    )
    add_before(
        anchor,
        "Distributed/cloud/AI product breadth is deferred. Relevant low-level "
        "kernel connections remain integrated where technically necessary.",
        "List Bullet",
    )
    add_before(
        anchor,
        "Do not auto-fill learner confidence, hours, completion, or mastery. "
        "Those remain honest self-assessments.",
        "List Bullet",
    )
    marker = add_before(anchor, END)
    for run in marker.runs:
        run.font.hidden = True

    document.save(TRACKER)
    print(f"updated {TRACKER}")


if __name__ == "__main__":
    main()
